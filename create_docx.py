import docx
from datetime import datetime

doc = docx.Document()

doc.add_heading('Location Page Blank Screen Bug - Implementation Plan', 0)

doc.add_paragraph('Date: 2026-06-09')
doc.add_paragraph('Time: 09:28:32 EST')

doc.add_heading('Problem Statement', level=1)
doc.add_paragraph('When the location intelligence API (/api/location-intel) experiences a failure (e.g., a timeout or a 500 server error), the location page renders completely blank instead of displaying the appropriate error or timeout banner to the user.')

doc.add_heading('Root Cause Analysis', level=1)
p = doc.add_paragraph('1. In ')
p.add_run('src/lib/location/components/AddressAnalyzer.svelte').bold = True
p.add_run(', when liveIntel fails to fetch and returns null, the component logs a warning but skips populating compassScores. It either fails to emit the onScoresReady event or emits it with missing data.')

p = doc.add_paragraph('2. In ')
p.add_run('src/routes/app/location/+page.svelte').bold = True
p.add_run(', the UI relies on a $derived state called locationIQ. If the sixScores object passed from the analyzer is empty, locationIQ evaluates to 0.')

p = doc.add_paragraph('3. The main UI blocks (<div class="v3-hero"> and <div class="v3-split">) are guarded by conditions like ')
p.add_run('{#if fitIQ > 0 || locationIQ > 0}').italic = True
p.add_run('.')

doc.add_paragraph('4. Because the score evaluates to 0, Svelte hides all of the main content blocks. While there is an analysisError banner intended for this scenario, it relies on a scoringFailed flag that isn\'t being properly propagated from the AddressAnalyzer when the payload is entirely dropped.')

doc.add_heading('Proposed Solution', level=1)
p = doc.add_paragraph('1. ')
p.add_run('Fix AddressAnalyzer.svelte to explicitly emit a failure state:').bold = True
p.add_run('\nModify the else block (when liveIntel is null) and the error catch blocks in AddressAnalyzer.svelte to explicitly fire onScoresReady({ scoringTimedOut: true, compassScores: {}, compassComposite: 0 }). This ensures the parent component is notified that the search concluded but failed.')

p = doc.add_paragraph('2. ')
p.add_run('Update +page.svelte to ensure the Error Banner renders:').bold = True
p.add_run('\nEnsure that the <div class="analysis-error-banner"> is displayed properly when analysisError is set, making sure it isn\'t trapped inside a block that requires locationIQ > 0.')

doc.add_heading('Next Steps', level=1)
doc.add_paragraph('- Review and approve this plan.')
doc.add_paragraph('- Once approved, I will implement these changes in AddressAnalyzer.svelte and +page.svelte so the user receives clear visual feedback instead of a blank page.')

output_path = r'C:\sandbox_re_squared\jared_prod\techstack_revamp\Location_Page_Blank_Screen_Fix_2026-06-09_09-28-32.docx'
doc.save(output_path)
print("Docx saved to", output_path)
