import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import datetime

doc = docx.Document()

# ── Styles ──
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

def heading(text, level=1):
    h = doc.add_heading(text, level=level)
    h.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    return h

def bold_para(label, value):
    p = doc.add_paragraph()
    r = p.add_run(label)
    r.bold = True
    p.add_run(value)
    return p

def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    # Header row
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for para in hdr[i].paragraphs:
            for run in para.runs:
                run.bold = True
        tc = hdr[i]._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), '1F497D')
        shd.set(qn('w:color'), 'FFFFFF')
        shd.set(qn('w:val'), 'clear')
        tcPr.append(shd)
        for para in hdr[i].paragraphs:
            for run in para.runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    # Data rows
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = val
            if ri % 2 == 1:
                tc = cells[ci]._tc
                tcPr = tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:fill'), 'DCE6F1')
                shd.set(qn('w:val'), 'clear')
                tcPr.append(shd)
    return table

# ── Title Page ──
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = title.add_run('Snowflake → Cloud SQL Migration & Sync Architecture')
tr.bold = True
tr.font.size = Pt(20)
tr.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

doc.add_paragraph()
ts_para = doc.add_paragraph()
ts_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
ts_para.add_run(f'Date: {datetime.now().strftime("%B %d, %Y  |  %I:%M %p EST")}').italic = True

doc.add_paragraph()

# ── 1. Executive Summary ──
heading('1. Executive Summary')
doc.add_paragraph(
    'This document defines the architecture for replicating the Snowflake analytical warehouse '
    'data into a Google Cloud SQL (PostgreSQL) instance that is directly accessible by the '
    'SvelteKit application backend. The goal is to completely eliminate 15-second cold-start '
    'delays and 504 Gateway Timeout errors caused by querying Snowflake at request time. '
    'Cloud SQL will serve as the OLTP (Online Transaction Processing) read layer for the app, '
    'while Snowflake remains the OLAP (Online Analytical Processing) engine for data science, '
    'dbt models, and batch scoring.'
)

# ── 2. Problem Statement ──
heading('2. Problem Statement')
rows_problem = [
    ('Cold Start Latency', 'Snowflake COMPUTE_WH auto-suspends when idle. First query takes 10–15s to provision compute resources.'),
    ('Serverless Timeout', 'Netlify/Vercel serverless functions have a hard 26s execution limit. A 15s Snowflake cold start leaves only 11s for 21 API calls + AI layers.'),
    ('Blank Screen Bug', 'When the API times out, liveIntel returns null. The frontend has no score data, locationIQ = 0, and Svelte hides the entire UI.'),
    ('Snowflake SDK Overhead', 'snowflake-sdk adds significant bundle weight to the serverless function, increasing cold start time further.'),
]
add_table(['Problem', 'Impact'], rows_problem)
doc.add_paragraph()

# ── 3. Proposed Architecture ──
heading('3. Proposed Architecture')
doc.add_paragraph(
    'Snowflake remains the single source of truth for analytical scoring (dbt, batch jobs). '
    'A nightly sync pipeline replicates the final scored rows into Cloud SQL. '
    'The SvelteKit API reads exclusively from Cloud SQL at request time — no more Snowflake in the hot path.'
)
p = doc.add_paragraph()
p.add_run('Data Flow:').bold = True
doc.add_paragraph('Snowflake (OLAP / dbt scoring) → [Nightly Sync] → Cloud SQL Postgres (OLTP / app reads) → SvelteKit API → User Browser')

# ── 4. Table Mapping ──
heading('4. Table Mapping: Snowflake → Cloud SQL')

heading('4.1 fct_location_scores', level=2)
doc.add_paragraph('Core location intelligence scores — the primary table consumed by the app.')
rows_scores = [
    ('LOCATION_ID', 'TEXT', 'location_id', 'TEXT PRIMARY KEY', 'Unique identifier per location'),
    ('LOCATION_NAME', 'TEXT', 'location_name', 'TEXT', 'Human-readable address/name'),
    ('BUSINESS_CATEGORY', 'TEXT', 'business_category', 'TEXT', 'e.g. "cafe", "gym"'),
    ('TOTAL_COMPETITORS', 'NUMBER', 'total_competitors', 'INTEGER', 'Competitor POI count'),
    ('TOTAL_DAILY_RIDERSHIP', 'NUMBER', 'total_daily_ridership', 'INTEGER', 'MTA subway ridership'),
    ('ESTIMATED_MORNING_PASSERSBY', 'NUMBER', 'estimated_morning_passersby', 'INTEGER', 'Foot traffic estimate'),
    ('MEDIAN_HOUSEHOLD_INCOME', 'NUMBER', 'median_household_income', 'INTEGER', 'Census ACS income'),
    ('CRIME_COUNT', 'NUMBER', 'crime_count', 'INTEGER', 'NYPD complaints in radius'),
    ('COMPLAINT_COUNT', 'NUMBER', 'complaint_count', 'INTEGER', '311 complaints in radius'),
    ('BUSINESS_LICENSE_COUNT', 'NUMBER', 'business_license_count', 'INTEGER', 'DCA licenses in radius'),
    ('MORNING_FOOT_TRAFFIC_SCORE', 'FLOAT', 'morning_foot_traffic_score', 'NUMERIC(5,2)', '0-100 foot traffic score'),
    ('COMPETITION_CONTEXT_SCORE', 'FLOAT', 'competition_context_score', 'NUMERIC(5,2)', '0-100 competition score'),
    ('DEMOGRAPHICS_FIT_SCORE', 'FLOAT', 'demographics_fit_score', 'NUMERIC(5,2)', '0-100 demographics score'),
    ('DAILY_RITUAL_DENSITY_SCORE', 'FLOAT', 'daily_ritual_density_score', 'NUMERIC(5,2)', '0-100 density score'),
    ('STREET_SIDE_SCORE', 'FLOAT', 'street_side_score', 'NUMERIC(5,2)', '0-100 street side score'),
    ('SAFETY_SCORE', 'FLOAT', 'safety_score', 'NUMERIC(5,2)', '0-100 safety score'),
    ('NEIGHBORHOOD_HEALTH_SCORE', 'FLOAT', 'neighborhood_health_score', 'NUMERIC(5,2)', '0-100 neighborhood score'),
    ('BUSINESS_SURVIVAL_SCORE', 'FLOAT', 'business_survival_score', 'NUMERIC(5,2)', '0-100 survival score'),
    ('FINAL_LOCATION_IQ', 'FLOAT', 'final_location_iq', 'NUMERIC(5,2)', 'Final 0-100 IQ score'),
    ('CATEGORY_PERCENTILE', 'FLOAT', 'category_percentile', 'NUMERIC(5,4)', 'Percentile rank 0.0-1.0'),
    ('SCORED_AT', 'TIMESTAMP_NTZ', 'scored_at', 'TIMESTAMPTZ', 'When Snowflake scored this row'),
    ('SYNCED_AT', '(derived)', 'synced_at', 'TIMESTAMPTZ DEFAULT now()', 'When Cloud SQL received this row'),
]
add_table(['Snowflake Column', 'Snowflake Type', 'Cloud SQL Column', 'Cloud SQL Type', 'Notes'], rows_scores)
doc.add_paragraph()

heading('4.2 intel_cache (existing Supabase table)', level=2)
doc.add_paragraph(
    'This table already exists in Supabase/Postgres and stores the raw 21-source API '
    'payloads (Census, WalkScore, Google Places, etc.). It does NOT need to be replicated '
    'from Snowflake — it is written directly by the SvelteKit API. It serves as the '
    'fallback data layer when a location has not yet been scored by Snowflake.'
)
rows_cache = [
    ('cache_key', 'TEXT PRIMARY KEY', 'Lat/lng/businessType composite key'),
    ('source', 'TEXT', 'Source label e.g. "location-report"'),
    ('data', 'JSONB', 'Full LocationIntelReport payload'),
    ('fetched_at', 'TIMESTAMPTZ', 'When the data was fetched'),
    ('ttl_ms', 'INTEGER', 'Time-to-live in milliseconds'),
]
add_table(['Column', 'Type', 'Description'], rows_cache)
doc.add_paragraph()

# ── 5. Sync Strategy ──
heading('5. Sync Strategy: Keeping Snowflake & Cloud SQL in Sync')

heading('5.1 Recommended: Google Cloud Scheduler + Python Script', level=2)
doc.add_paragraph(
    'A lightweight Python Cloud Run Job runs on a nightly schedule (e.g. 2:00 AM EST). '
    'It pulls all rows updated in Snowflake since the last sync and upserts them into Cloud SQL.'
)
doc.add_paragraph('Step-by-step:')
steps = [
    '1. Cloud Scheduler triggers the Cloud Run Job at 2:00 AM EST.',
    '2. The job connects to Snowflake and runs: SELECT * FROM fct_location_scores WHERE SCORED_AT > <last_synced_at>.',
    '3. For each row returned, the job performs an UPSERT into Cloud SQL using ON CONFLICT (location_id) DO UPDATE.',
    '4. The job logs the number of rows synced and stores the latest SCORED_AT as the next watermark.',
    '5. Done. The entire process takes under 60 seconds for typical data volumes.',
]
for s in steps:
    doc.add_paragraph(s, style='List Bullet')
doc.add_paragraph()

heading('5.2 Sync Frequency Options', level=2)
rows_sync = [
    ('Nightly (Recommended)', 'Location scores change daily based on new crime/inspection data.', 'Cloud Scheduler cron: 0 2 * * *', 'Low cost, simple'),
    ('Hourly', 'If Snowflake dbt models run hourly for high-frequency scoring.', 'Cloud Scheduler cron: 0 * * * *', 'Higher Snowflake credit cost'),
    ('On-Demand (Webhook)', 'Trigger sync immediately after a dbt model run completes.', 'dbt on-run-end hook → Cloud Run', 'Most real-time, most complex'),
    ('Real-time (Snowflake Streams)', 'Snowflake Streams capture row-level changes and emit events.', 'Snowflake Stream → Pub/Sub → Cloud SQL', 'Most powerful, highest cost'),
]
add_table(['Strategy', 'When to Use', 'Mechanism', 'Trade-off'], rows_sync)
doc.add_paragraph()

heading('5.3 Conflict / Staleness Handling', level=2)
conflicts = [
    ('Row not yet in Snowflake', 'New address searched but not yet scored.', 'Fall back to live intel_cache data. Return a "Scoring in progress" badge to the user.'),
    ('Stale row in Cloud SQL', 'Score is > 7 days old.', 'Add a scored_at column check in the API. If stale, trigger a background re-score request.'),
    ('Sync job failure', 'Cloud Run Job crashes mid-sync.', 'Use idempotent upserts. Re-running the job will safely re-process failed rows.'),
    ('Schema drift', 'New column added to Snowflake fct_location_scores.', 'Versioned migration scripts (Drizzle) must be applied to Cloud SQL before the next sync.'),
]
add_table(['Scenario', 'Cause', 'Resolution'], conflicts)
doc.add_paragraph()

# ── 6. Code Changes Required ──
heading('6. Code Changes Required in SvelteKit')
changes = [
    ('src/lib/snowflake.ts', 'DEPRECATE', 'Remove direct Snowflake SDK calls from the hot API path.'),
    ('src/lib/golden-record.ts', 'NEW', 'New module. Reads fct_location_scores from Cloud SQL via Drizzle ORM.'),
    ('src/lib/db/schema.ts', 'MODIFY', 'Add Drizzle table definition for location_scores.'),
    ('src/routes/api/location-iq/+server.ts', 'MODIFY', 'Replace getGoldenRecord() (Snowflake) with getLocationScore() (Cloud SQL).'),
    ('sync/snowflake_to_cloudsql.py', 'NEW', 'Python sync script for Cloud Run Job.'),
    ('sync/cloudbuild.yaml', 'NEW', 'Cloud Build config to containerize and deploy the sync job.'),
]
add_table(['File', 'Action', 'Description'], changes)
doc.add_paragraph()

# ── 7. Benefits Summary ──
heading('7. Expected Benefits After Migration')
benefits = [
    ('API Response Time', '10–25 seconds (cold start)', '< 100ms (always warm Postgres)'),
    ('Blank Screen Bug', 'Occurs on every cold start', 'Eliminated entirely'),
    ('504 Timeout Errors', 'Frequent during low-traffic periods', 'Zero'),
    ('Snowflake SDK Dependency', 'Required in serverless function', 'Removed — smaller bundle'),
    ('Data Freshness', 'Real-time (but unreliable)', 'Nightly (stable and predictable)'),
    ('Cost', 'Snowflake credits on every user query', 'Snowflake credits only on nightly batch'),
]
add_table(['Metric', 'Before', 'After'], benefits)

doc.add_paragraph()
doc.add_paragraph('End of Document', style='Normal').runs[0].italic = True

# ── Save ──
timestamp = datetime.now().strftime('%Y-%m-%d_%H-%M-%S')
output_path = fr'C:\sandbox_re_squared\jared_prod\techstack_revamp\Snowflake-CloudSql_{timestamp}.docx'
doc.save(output_path)
print(f"Saved: {output_path}")
