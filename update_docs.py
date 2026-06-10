import pandas as pd
from docx import Document
import os

excel_file = 'supabase-drizzle-ai-endpoint-mapping.05312026.xlsx'

# 1. Update Excel
df = pd.read_excel(excel_file)
changed_files = [
    'audit.ts',
    'budget-guard.ts',
    'llm-cache.ts',
    'telemetry.ts',
    'compare-v2.ts',
    'survival-rate.ts',
    'bundle-builder.ts',
    'data-quality-gate.ts',
    'capture.ts'
]

def update_status(row):
    endpoint = str(row['Endpoint']).lower()
    for cf in changed_files:
        if cf in endpoint:
            return 'Completed'
    return row['Status']

df['Status'] = df.apply(update_status, axis=1)
df.to_excel(excel_file, index=False)

# 2. Create Word Document
doc = Document()
doc.add_heading('Supabase Migration & Cleanup Summary', 0)

doc.add_heading('Overview', level=1)
doc.add_paragraph(
    "This document summarizes the final cleanup of all Supabase dependencies from the application. "
    "All data access has been successfully migrated to use the Drizzle ORM via $lib/db-server."
)

doc.add_heading('Files Refactored', level=1)
ul = doc.add_paragraph(style='List Bullet')
ul.add_run('src/lib/audit.ts').bold = True
ul.add_run(' - Swapped getServiceSupabase for db.execute(sql...)')

ul2 = doc.add_paragraph(style='List Bullet')
ul2.add_run('src/lib/ai/budget-guard.ts').bold = True
ul2.add_run(' - Swapped getServiceSupabase for db.execute(sql...)')

ul3 = doc.add_paragraph(style='List Bullet')
ul3.add_run('src/lib/ai/llm-cache.ts').bold = True
ul3.add_run(' - Swapped getServiceSupabase for db.execute(sql...)')

ul4 = doc.add_paragraph(style='List Bullet')
ul4.add_run('src/lib/ai/telemetry.ts').bold = True
ul4.add_run(' - Swapped getServiceSupabase for db.execute(sql...)')

ul5 = doc.add_paragraph(style='List Bullet')
ul5.add_run('src/lib/intel/compare-v2.ts').bold = True
ul5.add_run(' - Swapped getServiceSupabase for db.execute(sql...)')

ul6 = doc.add_paragraph(style='List Bullet')
ul6.add_run('src/lib/intel/survival-rate.ts').bold = True
ul6.add_run(' - Swapped getServiceSupabase for db.execute(sql...)')

ul7 = doc.add_paragraph(style='List Bullet')
ul7.add_run('src/lib/intel/scoring/bundle-builder.ts').bold = True
ul7.add_run(' - Swapped getServiceSupabase for db.execute(sql...)')

ul8 = doc.add_paragraph(style='List Bullet')
ul8.add_run('src/lib/intel/data-quality-gate.ts').bold = True
ul8.add_run(' - Swapped getServiceSupabase for db.execute(sql...)')

ul9 = doc.add_paragraph(style='List Bullet')
ul9.add_run('src/lib/analytics/capture.ts').bold = True
ul9.add_run(' - Removed unused Supabase telemetry functions')

doc.add_heading('Build Verification', level=1)
doc.add_paragraph(
    "After removing all imports of '@supabase/supabase-js' and deleting the unused local supabase clients, "
    "the application successfully compiled using 'npm run build' with zero errors. All integration points have been validated."
)

doc.save('Supabase_Cleanup_Summary.docx')
print("Docs updated successfully.")
